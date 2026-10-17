"""Run spaCy AND the Hugging Face transformer on the answer-key paragraphs and write one Excel file
to share with the team (task 4.3 -> hand-off to Yeshasvini for 4.1 / 4.4).

Usage:   python compare_models.py
Output:  results/NER_Model_Comparison.xlsx   (also prints the scores)

If a model cannot be loaded (for example no internet for the first transformer download),
it is skipped with a message and the file is still written for the other model.
"""
import argparse
import os
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from evaluate import evaluate, norm, read_csv
from ner_engine import ENTITY_TYPES, MODELS, NERError, extract

HDR_FILL = PatternFill("solid", fgColor="1F4E79")
HDR_FONT = Font(bold=True, color="FFFFFF")
ERROR_LABEL = {
    "false_positive": "Wrong extra (model found it, key does not)",
    "false_negative": "Missed (in key, model did not get it)",
}


def prf(c):
    p = c["tp"] / (c["tp"] + c["fp"]) if c["tp"] + c["fp"] else 0.0
    r = c["tp"] / (c["tp"] + c["fn"]) if c["tp"] + c["fn"] else 0.0
    f = 2 * p * r / (p + r) if p + r else 0.0
    return round(p, 2), round(r, 2), round(f, 2)


def write_sheet(ws, headers, rows, widths):
    ws.append(headers)
    for cell in ws[1]:
        cell.fill, cell.font = HDR_FILL, HDR_FONT
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--samples", default="data/samples.csv")
    ap.add_argument("--gold", default="data/gold_entities.csv")
    ap.add_argument("--out", default="results/NER_Model_Comparison.xlsx")
    args = ap.parse_args()

    samples, gold_rows = read_csv(args.samples), read_csv(args.gold)
    gold = defaultdict(set)
    for r in gold_rows:
        gold[r["sample_id"].strip()].add((norm(r["entity"]), r["type"].strip().upper()))
    scored = [s for s in samples if s["sample_id"].strip() in gold]

    results, preds, skipped = {}, {}, []
    for model in ("spacy", "transformer"):
        try:
            counts, errors = evaluate(samples, gold_rows, model)
            preds[model] = {
                s["sample_id"].strip(): {(norm(e["text"]), e["label"]) for e in extract(s["text"], model)}
                for s in scored
            }
            results[model] = (counts, errors)
            print(f"[ok] {MODELS[model]}")
        except NERError as err:
            skipped.append(model)
            print(f"[skipped] {MODELS[model]}: {err}")

    if not results:
        raise SystemExit("No model could run. Install the requirements first (see README.md).")
    names = {m: MODELS[m] for m in results}

    wb = Workbook()
    # --- Summary ---
    ws = wb.active
    ws.title = "Summary"
    rows = []
    for typ in ENTITY_TYPES + ["ALL"]:
        row = [typ]
        for m, (counts, _) in results.items():
            c = {k: sum(counts[t][k] for t in ENTITY_TYPES) for k in ("tp", "fp", "fn")} if typ == "ALL" else counts[typ]
            p, r, f = prf(c)
            row += [c["tp"], c["fp"], c["fn"], p, r, f]
        rows.append(row)
    headers = ["Type"]
    for m in results:
        headers += [f"{n} {names[m]}" for n in ("TP:", "FP:", "FN:", "Precision:", "Recall:", "F1:")]
    write_sheet(ws, headers, rows, [10] + [16] * (6 * len(results)))
    ws.append([])
    ws.append([f"Scored on {len(scored)} paragraphs with an answer key: " + ", ".join(s["sample_id"] for s in scored)])
    ws.append(["TP = correct, FP = wrong extra, FN = missed. An entity is correct only if text AND type match."])
    for m in skipped:
        ws.append([f"NOT RUN: {MODELS[m]} (could not be loaded)"])

    # --- Per paragraph ---
    ws2 = wb.create_sheet("Per paragraph")
    fmt = lambda S: "\n".join(f"{e} ({t})" for e, t in sorted(S))
    rows = []
    for s in scored:
        sid = s["sample_id"].strip()
        row = [sid, fmt(gold[sid])]
        for m in results:
            row += [fmt(preds[m][sid]), fmt(preds[m][sid] - gold[sid]), fmt(gold[sid] - preds[m][sid])]
        rows.append(row)
    headers = ["Paragraph", "Answer key"]
    for m in results:
        headers += [f"{names[m]}: found", f"{names[m]}: wrong extra", f"{names[m]}: missed"]
    write_sheet(ws2, headers, rows, [11, 30] + [28] * (3 * len(results)))

    # --- Mistakes (for task 4.4) ---
    ws3 = wb.create_sheet("Mistakes (fill in cause)")
    rows = [
        [e["sample_id"], names[m], ERROR_LABEL[e["error"]], e["entity"], e["type"], ""]
        for m, (_, errors) in results.items()
        for e in errors
    ]
    write_sheet(ws3, ["Paragraph", "Model", "Error type", "Entity (lower-case)", "Type", "Cause (to fill in)"],
                rows, [11, 24, 40, 24, 9, 45])

    os.makedirs(os.path.dirname(args.out) or ".", exist_ok=True)
    wb.save(args.out)
    print(f"\nSaved {args.out}")
    for m, (counts, _) in results.items():
        total = {k: sum(counts[t][k] for t in ENTITY_TYPES) for k in ("tp", "fp", "fn")}
        p, r, f = prf(total)
        print(f"{names[m]:<26} precision {p:.2f}  recall {r:.2f}  F1 {f:.2f}")


if __name__ == "__main__":
    main()
